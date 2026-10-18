from io import BytesIO
import sqlite3
from pathlib import Path

import fitz
import pytest
from docx import Document
from fastapi.testclient import TestClient

from backend.app import documents as documents_api
from backend.app.chunker import chunk_document, chunk_text
from backend.app.document_processor import process_document
from backend.app.main import app
from backend.app.retriever import EmbeddingError, IndexingResult


class FakeUploadRetriever:
    def __init__(self, fail: bool = False) -> None:
        self.calls: list[tuple[str, str, list[object]]] = []
        self.fail = fail

    def index_document(self, document_id: str, filename: str, chunks: list[object]) -> IndexingResult:
        self.calls.append((document_id, filename, chunks))
        if self.fail:
            raise EmbeddingError("embedding model unavailable")
        return IndexingResult(
            document_id=document_id,
            status="indexed" if chunks else "empty",
            chunk_count=len(chunks),
        )


class FakeUploadService:
    def __init__(self) -> None:
        self.retriever = FakeUploadRetriever()


@pytest.fixture
def client(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    auth_headers: dict[str, str],
) -> TestClient:
    monkeypatch.setattr(documents_api, "UPLOAD_DIR", tmp_path)
    catalog_path = tmp_path.parent / f"{tmp_path.name}-catalog.sqlite3"
    monkeypatch.setattr(documents_api, "DOCUMENT_CATALOG_PATH", catalog_path)
    documents_api._initialize_document_catalog()
    with sqlite3.connect(str(catalog_path)) as connection:
        connection.execute("DELETE FROM documents")
    fake_service = FakeUploadService()
    monkeypatch.setattr(documents_api, "get_default_rag_service", lambda: fake_service)
    return TestClient(app, headers=auth_headers)


def make_pdf(page_texts: list[str]) -> bytes:
    pdf = fitz.open()
    for page_text in page_texts:
        page = pdf.new_page()
        if page_text:
            page.insert_text((72, 72), page_text)
    content = pdf.tobytes()
    pdf.close()
    return content


def make_docx() -> bytes:
    document = Document()
    document.add_paragraph("Research notes")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Term"
    table.cell(0, 1).text = "Meaning"
    table.add_row().cells[0].text = "RAG"
    table.rows[1].cells[1].text = "Retrieval augmented generation"
    content = BytesIO()
    document.save(content)
    return content.getvalue()


def test_upload_pdf_extracts_multiple_pages(client: TestClient) -> None:
    response = client.post(
        "/api/v1/documents/upload",
        files={"file": ("report.pdf", make_pdf(["First page", "Second page"]), "application/pdf")},
    )

    assert response.status_code == 200
    result = response.json()
    assert result["file_type"] == "pdf"
    assert result["extraction_status"] == "success"
    assert [page["page_number"] for page in result["pages"]] == [1, 2]
    assert "First page" in result["extracted_text"]
    assert "Second page" in result["extracted_text"]


def test_upload_docx_extracts_paragraphs_and_tables(client: TestClient) -> None:
    response = client.post(
        "/api/v1/documents/upload",
        files={"file": ("notes.docx", make_docx(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )

    assert response.status_code == 200
    result = response.json()
    assert "Research notes" in result["extracted_text"]
    assert "Term | Meaning" in result["extracted_text"]
    assert "Retrieval augmented generation" in result["extracted_text"]
    assert result["metadata"]["table_count"] == 1


def test_upload_csv_preserves_columns_and_row_context(client: TestClient) -> None:
    response = client.post(
        "/api/v1/documents/upload",
        files={"file": ("metrics.csv", b"department,revenue\nResearch,42\n", "text/csv")},
    )

    assert response.status_code == 200
    result = response.json()
    assert result["metadata"] == {
        "column_names": ["department", "revenue"],
        "row_count": 1,
    }
    assert result["rows"] == [{"department": "Research", "revenue": "42"}]
    assert "Row 1: department=Research | revenue=42" in result["extracted_text"]
    assert result["indexing_status"] == "indexed"
    assert result["chunk_count"] == 1


def test_empty_pdf_is_reported_as_empty(client: TestClient) -> None:
    response = client.post(
        "/api/v1/documents/upload",
        files={"file": ("blank.pdf", make_pdf([""]), "application/pdf")},
    )

    assert response.status_code == 200
    assert response.json()["extraction_status"] == "empty"


def test_corrupted_pdf_is_rejected_and_removed(client: TestClient, tmp_path: Path) -> None:
    response = client.post(
        "/api/v1/documents/upload",
        files={"file": ("broken.pdf", b"%PDF-1.7\nnot a valid PDF", "application/pdf")},
    )

    assert response.status_code == 422
    assert response.json()["detail"] == "The uploaded document is invalid or could not be extracted."
    assert list(tmp_path.iterdir()) == []


def test_unsupported_file_type_is_rejected(client: TestClient, tmp_path: Path) -> None:
    response = client.post(
        "/api/v1/documents/upload",
        files={"file": ("data.xlsx", b"not an allowed format", "application/octet-stream")},
    )

    assert response.status_code == 415
    assert list(tmp_path.iterdir()) == []


def test_oversized_upload_is_rejected(client: TestClient, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(documents_api, "MAX_UPLOAD_SIZE_BYTES", 16)
    response = client.post(
        "/api/v1/documents/upload",
        files={"file": ("large.csv", b"x" * 17, "text/csv")},
    )

    assert response.status_code == 413


def test_malformed_csv_is_rejected(client: TestClient, tmp_path: Path) -> None:
    response = client.post(
        "/api/v1/documents/upload",
        files={"file": ("invalid.csv", b"\xff\xfe", "text/csv")},
    )

    assert response.status_code == 422
    assert list(tmp_path.iterdir()) == []


def test_inconsistent_csv_row_width_is_rejected(client: TestClient, tmp_path: Path) -> None:
    response = client.post(
        "/api/v1/documents/upload",
        files={"file": ("invalid.csv", b"name,value\nAda,42,extra\n", "text/csv")},
    )

    assert response.status_code == 422
    assert list(tmp_path.iterdir()) == []


def test_upload_filename_cannot_choose_storage_path(client: TestClient, tmp_path: Path) -> None:
    response = client.post(
        "/api/v1/documents/upload",
        files={"file": ("../../escape.csv", b"name\nvalue\n", "text/csv")},
    )

    assert response.status_code == 200
    assert response.json()["original_filename"] == "escape.csv"
    stored_files = list(tmp_path.iterdir())
    assert len(stored_files) == 1
    assert stored_files[0].name != "escape.csv"


def test_upload_reports_index_failure_without_claiming_searchability(
    client: TestClient,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    failed_service = FakeUploadService()
    failed_service.retriever.fail = True
    monkeypatch.setattr(
        documents_api,
        "get_default_rag_service",
        lambda: failed_service,
    )

    response = client.post(
        "/api/v1/documents/upload",
        files={"file": ("policy.csv", b"policy\nannual leave\n", "text/csv")},
    )

    assert response.status_code == 200
    assert response.json()["extraction_status"] == "success"
    assert response.json()["indexing_status"] == "failed"
    assert response.json()["indexing_error"].endswith("not searchable.")


def test_document_listing_does_not_return_extracted_content(client: TestClient) -> None:
    client.post(
        "/api/v1/documents/upload",
        files={"file": ("metrics.csv", b"department\nResearch\n", "text/csv")},
    )

    result = client.get("/api/v1/documents")

    assert result.status_code == 200
    listed_document = result.json()[0]
    assert listed_document["original_filename"] == "metrics.csv"
    assert "extracted_text" not in listed_document
    assert "rows" not in listed_document


def test_chunk_text_creates_overlapping_chunks_with_metadata() -> None:
    chunks = chunk_text(
        "abcdefghijklmnopqrst",
        "report.pdf",
        chunk_size=10,
        overlap=3,
        page_number=2,
    )

    assert [chunk.text for chunk in chunks] == ["abcdefghij", "hijklmnopq", "opqrst"]
    assert [chunk.metadata["chunk_index"] for chunk in chunks] == [0, 1, 2]
    assert all(chunk.metadata["page_number"] == 2 for chunk in chunks)


def test_chunk_text_returns_no_chunks_for_empty_text() -> None:
    assert chunk_text("  \n", "empty.docx") == []


def test_chunk_document_preserves_page_order_and_numbers(tmp_path: Path) -> None:
    path = tmp_path / "ordered.pdf"
    path.write_bytes(make_pdf(["A" * 12, "B" * 12]))
    document = process_document(path, "ordered.pdf")

    chunks = chunk_document(document, chunk_size=8, overlap=2)

    assert [chunk.metadata["page_number"] for chunk in chunks] == [1, 1, 2, 2]
    assert [chunk.metadata["chunk_index"] for chunk in chunks] == [0, 1, 2, 3]