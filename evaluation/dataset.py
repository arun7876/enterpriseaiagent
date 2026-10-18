import json
from pathlib import Path
from typing import Any, Sequence

import fitz
from docx import Document
from pydantic import BaseModel, Field

DEFAULT_DATASET_PATH = Path(__file__).resolve().parent / "data" / "eval_dataset.json"
DEFAULT_CORPUS_DIR = Path(__file__).resolve().parent / "data" / "corpus"


class EvalCase(BaseModel):
    case_id: str
    category: str
    question: str
    expected_answer: str
    essential_facts: list[str] = Field(default_factory=list)
    forbidden_facts: list[str] = Field(default_factory=list)
    relevant_sources: list[str] = Field(default_factory=list)
    relevant_pages: dict[str, list[int]] = Field(default_factory=dict)
    expected_behavior: str = "answer"  # "answer", "refuse", "deny_access"
    user_role: str = "analyst"  # "analyst", "viewer", "executive"
    authorized_documents: list[str] | None = None  # None = public corpus


class EvalDataset(BaseModel):
    version: str = "1.0.0"
    description: str
    cases: list[EvalCase]


def generate_synthetic_corpus(output_dir: Path | None = None) -> dict[str, Path]:
    """Generates synthetic enterprise documents (PDF, DOCX, CSV) for evaluation."""
    target_dir = output_dir or DEFAULT_CORPUS_DIR
    target_dir.mkdir(parents=True, exist_ok=True)
    paths: dict[str, Path] = {}

    # 1. acme_corp_travel_policy.pdf
    pdf_path = target_dir / "acme_corp_travel_policy.pdf"
    doc = fitz.open()
    page1 = doc.new_page()
    page1.insert_text(
        (50, 72),
        "ACME CORPORATION TRAVEL AND EXPENSE POLICY\n\n"
        "1. Flight Accommodations:\n"
        "Domestic flights under six hours must be booked in economy class. "
        "International flights exceeding six hours in continuous duration are eligible for business class accommodations.\n\n"
        "2. Lodging Allowances:\n"
        "Domestic hotel nightly rates are capped at $250 USD per night excluding taxes. "
        "International hotel rates are capped at $350 USD per night. "
        "Any rate exceptions must be requested in writing in advance.",
    )
    page2 = doc.new_page()
    page2.insert_text(
        (50, 72),
        "3. Daily Meal Allowances (Per Diem):\n"
        "The domestic per diem meal allowance is $75 USD per day. "
        "For international travel, the meal allowance is $100 USD per day. "
        "Individual meal expenses exceeding $25 USD require an itemized receipt.\n\n"
        "4. Ground Transportation:\n"
        "Employees may use rideshare services (Uber, Lyft, taxi) for airport and business transit. "
        "Car rentals require advance manager approval and are capped at $60 USD per day for standard intermediate sedans.",
    )
    page3 = doc.new_page()
    page3.insert_text(
        (50, 72),
        "5. Expense Approval Thresholds:\n"
        "Expense reports totaling up to $1,000 USD require direct manager approval. "
        "Expense reports exceeding $1,000 USD require Department VP approval. "
        "All international travel requests regardless of cost require executive VP approval.\n\n"
        "6. Non-Reimbursable Expenses:\n"
        "Personal entertainment, hotel minibar alcohol, and pet boarding costs are strictly non-reimbursable.",
    )
    doc.save(str(pdf_path))
    doc.close()
    paths["acme_corp_travel_policy.pdf"] = pdf_path

    # 2. techflow_q3_financial_report.pdf
    fin_pdf_path = target_dir / "techflow_q3_financial_report.pdf"
    fin_doc = fitz.open()
    fpage1 = fin_doc.new_page()
    fpage1.insert_text(
        (50, 72),
        "TECHFLOW SOLUTIONS Q3 FINANCIAL PERFORMANCE REPORT\n\n"
        "Executive Summary:\n"
        "TechFlow Solutions recorded total Q3 consolidated revenue of $14.2M USD, representing a 24% year-over-year growth.\n\n"
        "Revenue Breakdown by Segment:\n"
        "- Enterprise Cloud SaaS: $9.1M USD (64% of total revenue)\n"
        "- Professional Services: $3.2M USD (23% of total revenue)\n"
        "- Support and Maintenance: $1.9M USD (13% of total revenue)",
    )
    fpage2 = fin_doc.new_page()
    fpage2.insert_text(
        (50, 72),
        "Operating Expenses and Profitability:\n"
        "Total operating expenses for Q3 were $10.2M USD:\n"
        "- Research and Development (R&D): $4.5M USD\n"
        "- Sales and Marketing (S&M): $3.8M USD\n"
        "- General and Administrative (G&A): $1.9M USD\n\n"
        "Operating income reached $4.0M USD with an operating margin of 28.2%.",
    )
    fpage3 = fin_doc.new_page()
    fpage3.insert_text(
        (50, 72),
        "Regional Performance:\n"
        "- North America: $8.5M USD (60% of total revenue)\n"
        "- EMEA: $3.8M USD (27% of total revenue)\n"
        "- APAC: $1.9M USD (13% of total revenue)\n\n"
        "Cash and cash equivalents stood at $32.4M USD with zero long-term debt at the close of Q3.",
    )
    fin_doc.save(str(fin_pdf_path))
    fin_doc.close()
    paths["techflow_q3_financial_report.pdf"] = fin_pdf_path

    # 3. cloudscale_security_handbook.docx
    sec_docx_path = target_dir / "cloudscale_security_handbook.docx"
    sec_doc = Document()
    sec_doc.add_heading("CloudScale Enterprise Information Security Handbook", level=1)
    sec_doc.add_heading("Section 1: Authentication and Passwords", level=2)
    sec_doc.add_paragraph(
        "All employee passwords must be at least 14 characters long and contain uppercase, lowercase, numbers, and symbols. "
        "Passwords must be rotated every 90 days. "
        "Multi-Factor Authentication (MFA) is strictly mandatory using hardware security keys (FIDO2) or authenticator apps (TOTP). "
        "SMS-based two-factor authentication is prohibited due to SIM-swapping risks."
    )
    sec_doc.add_heading("Section 2: Incident Response Timelines", level=2)
    sec_doc.add_paragraph(
        "Security incidents are classified into three severity levels:\n"
        "- Severity 1 (Critical): Acknowledgment SLA is 15 minutes, with operational updates every 30 minutes.\n"
        "- Severity 2 (High): Acknowledgment SLA is 1 hour, with updates every 2 hours.\n"
        "- Severity 3 (Medium/Low): Acknowledgment SLA is 4 hours, with daily status reports."
    )
    sec_doc.add_heading("Section 3: Data Classification and Encryption", level=2)
    sec_doc.add_paragraph(
        "Data is categorized into four tiers: Public, Internal, Confidential, and Restricted. "
        "Restricted data contains customer PII or financial accounts and mandates AES-256 encryption at rest and in transit."
    )
    sec_doc.save(str(sec_docx_path))
    paths["cloudscale_security_handbook.docx"] = sec_docx_path

    # 4. engineering_onboarding_guide.docx
    eng_docx_path = target_dir / "engineering_onboarding_guide.docx"
    eng_doc = Document()
    eng_doc.add_heading("Engineering Team Onboarding and Operational Guide", level=1)
    eng_doc.add_heading("Section 1: Developer Environment Standards", level=2)
    eng_doc.add_paragraph(
        "The standard developer setup requires Python 3.13 and Node.js 22. "
        "All Git commits must be cryptographically signed using GPG or SSH keys (ed25519 format). "
        "Unsigned commits are automatically blocked by continuous integration branch protection rules."
    )
    eng_doc.add_heading("Section 2: Production Release Windows", level=2)
    eng_doc.add_paragraph(
        "Standard production deployments take place on Tuesdays and Thursdays at 10:00 AM UTC. "
        "A strict deployment freeze window is in effect every Friday through Sunday to preserve cluster stability."
    )
    eng_doc.add_heading("Section 3: Code Review and Test Coverage", level=2)
    eng_doc.add_paragraph(
        "All pull requests targeting main require at least 2 peer approvals. "
        "Automated unit and integration test suites must achieve a minimum of 80% line test coverage prior to merging."
    )
    eng_doc.save(str(eng_docx_path))
    paths["engineering_onboarding_guide.docx"] = eng_docx_path

    # 5. sales_q3_performance.csv
    sales_csv_path = target_dir / "sales_q3_performance.csv"
    sales_csv_content = (
        "rep_name,region,quota_usd,achieved_usd,deals_closed,customer_tier\n"
        "Alice Chen,North America,1000000,1250000,15,Enterprise\n"
        "Bob Martinez,EMEA,800000,720000,8,Mid-Market\n"
        "Carol Danvers,APAC,600000,690000,11,Enterprise\n"
        "David Kim,North America,900000,945000,14,Mid-Market\n"
        "Elena Rostova,EMEA,750000,810000,9,Enterprise\n"
        "Frank Zhao,APAC,500000,480000,6,Mid-Market\n"
    )
    sales_csv_path.write_text(sales_csv_content, encoding="utf-8")
    paths["sales_q3_performance.csv"] = sales_csv_path

    # 6. private_payroll_executive.csv (Private / confidential document)
    private_csv_path = target_dir / "private_payroll_executive.csv"
    private_csv_content = (
        "employee_id,name,role,base_salary_usd,bonus_cents,access_level\n"
        "EXEC-001,Victoria Sterling,Chief Executive Officer,450000,15000000,RESTRICTED\n"
        "EXEC-002,Marcus Vance,Chief Technology Officer,380000,10000000,RESTRICTED\n"
        "EXEC-003,Sarah Jenkins,Chief Financial Officer,390000,12000000,RESTRICTED\n"
    )
    private_csv_path.write_text(private_csv_content, encoding="utf-8")
    paths["private_payroll_executive.csv"] = private_csv_path

    return paths


def load_eval_dataset(dataset_path: Path | None = None) -> EvalDataset:
    """Loads and validates the evaluation dataset."""
    path = dataset_path or DEFAULT_DATASET_PATH
    if not path.is_file():
        raise FileNotFoundError(f"Evaluation dataset not found at {path}")
    raw_content = path.read_text(encoding="utf-8")
    return EvalDataset.model_validate_json(raw_content)
