import csv
from pathlib import Path
from typing import Any, Literal
import zipfile

import fitz
import pandas as pd
from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from pydantic import BaseModel, Field


class DocumentProcessingError(ValueError):
    """Raised when a supported document cannot be extracted safely."""


class UnsupportedDocumentTypeError(ValueError):
    """Raised when a document does not use a supported file extension."""


class ExtractedPage(BaseModel):
    page_number: int
    text: str


class ProcessedDocument(BaseModel):
    original_filename: str
    file_type: Literal["pdf", "docx", "csv"]
    extracted_text: str
    extraction_status: Literal["success", "empty"]
    pages: list[ExtractedPage] = Field(default_factory=list)
    rows: list[dict[str, str]] = Field(default_factory=list)
    metadata: dict[str, Any] = Field(default_factory=dict)


def extract_pdf(file_path: str | Path) -> list[ExtractedPage]:
    path = Path(file_path)
    try:
        with path.open("rb") as pdf_file:
            header = pdf_file.read(1024)
        if b"%PDF-" not in header:
            raise DocumentProcessingError("The file is not a valid PDF document.")

        with fitz.open(str(path)) as pdf_document:
            if pdf_document.is_encrypted:
                raise DocumentProcessingError("Encrypted PDF documents are not supported.")
            return [
                ExtractedPage(page_number=page_number, text=page.get_text("text").strip())
                for page_number, page in enumerate(pdf_document, start=1)
            ]
    except DocumentProcessingError:
        raise
    except Exception as exc:
        raise DocumentProcessingError("The PDF document could not be extracted.") from exc


def extract_docx(file_path: str | Path) -> tuple[str, dict[str, int]]:
    path = Path(file_path)
    try:
        with zipfile.ZipFile(path) as docx_archive:
            required_parts = {"[Content_Types].xml", "word/document.xml"}
            if not required_parts.issubset(docx_archive.namelist()):
                raise DocumentProcessingError("The file is not a valid Word document.")

        document = Document(str(path))
        content_parts: list[str] = []
        paragraph_count = 0
        table_count = 0

        for child in document.element.body.iterchildren():
            if child.tag == qn("w:p"):
                paragraph_count += 1
                paragraph_text = Paragraph(child, document).text.strip()
                if paragraph_text:
                    content_parts.append(paragraph_text)
            elif child.tag == qn("w:tbl"):
                table_count += 1
                table = Table(child, document)
                table_rows = [
                    " | ".join(cell.text.strip() for cell in row.cells)
                    for row in table.rows
                ]
                if any(table_rows):
                    content_parts.append(
                        f"Table {table_count}:\n" + "\n".join(table_rows)
                    )

        return "\n\n".join(content_parts), {
            "paragraph_count": paragraph_count,
            "table_count": table_count,
        }
    except DocumentProcessingError:
        raise
    except Exception as exc:
        raise DocumentProcessingError("The Word document could not be extracted.") from exc


def _validate_csv_structure(file_path: str | Path) -> None:
    try:
        with Path(file_path).open("r", encoding="utf-8-sig", newline="") as csv_file:
            reader = csv.reader(csv_file, strict=True)
            header = next((row for row in reader if row), None)
            if header is None:
                return
            expected_columns = len(header)
            if any(len(row) != expected_columns for row in reader if row):
                raise DocumentProcessingError("The CSV document has inconsistent row widths.")
    except DocumentProcessingError:
        raise
    except (OSError, UnicodeDecodeError, csv.Error) as exc:
        raise DocumentProcessingError("The CSV document could not be parsed.") from exc


def extract_csv(file_path: str | Path) -> tuple[str, list[dict[str, str]], dict[str, Any]]:
    try:
        _validate_csv_structure(file_path)
        dataframe = pd.read_csv(
            file_path,
            dtype=str,
            keep_default_na=False,
            encoding="utf-8-sig",
            on_bad_lines="error",
        )
    except pd.errors.EmptyDataError:
        return "", [], {"column_names": [], "row_count": 0}
    except Exception as exc:
        raise DocumentProcessingError("The CSV document could not be parsed.") from exc

    column_names = [str(column) for column in dataframe.columns]
    rows = [
        {column: str(value) for column, value in row.items()}
        for row in dataframe.to_dict(orient="records")
    ]
    text_parts = [f"Columns: {', '.join(column_names)}"] if column_names else []
    text_parts.extend(
        "Row " + str(row_number) + ": " + " | ".join(
            f"{column}={row.get(column, '')}" for column in column_names
        )
        for row_number, row in enumerate(rows, start=1)
    )
    return "\n".join(text_parts), rows, {
        "column_names": column_names,
        "row_count": len(rows),
    }


def process_document(file_path: str | Path, filename: str) -> ProcessedDocument:
    normalized_filename = filename.replace("\\", "/").rsplit("/", maxsplit=1)[-1]
    file_type = Path(normalized_filename).suffix.lower().lstrip(".")

    if file_type == "pdf":
        pages = extract_pdf(file_path)
        extracted_text = "\n\n".join(
            f"[Page {page.page_number}]\n{page.text}" for page in pages if page.text
        )
        metadata: dict[str, Any] = {"page_count": len(pages)}
        rows: list[dict[str, str]] = []
    elif file_type == "docx":
        extracted_text, docx_metadata = extract_docx(file_path)
        pages = []
        metadata = docx_metadata
        rows = []
    elif file_type == "csv":
        extracted_text, rows, metadata = extract_csv(file_path)
        pages = []
    else:
        raise UnsupportedDocumentTypeError("Only PDF, DOCX, and CSV files are supported.")

    return ProcessedDocument(
        original_filename=normalized_filename,
        file_type=file_type,  # type: ignore[arg-type]
        extracted_text=extracted_text,
        extraction_status="success" if extracted_text.strip() else "empty",
        pages=pages,
        rows=rows,
        metadata=metadata,
    )